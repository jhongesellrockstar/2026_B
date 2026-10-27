from pathlib import Path
from PIL import Image
from docx import Document
from docx.shared import Cm, Pt, Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

w=Path(__file__).parent
for n,box in [(1,(8,104,385,365)),(2,(8,104,385,680)),(3,(8,240,385,418)),(4,(8,331,385,520))]:
    Image.open(w/f'P{n}raw.jpg').crop(box).save(w/f'P{n}.png')
d=Document(r'C:\Users\ACER\Downloads\Solucion_Analitica_Examen_Parcial_2026_2_Torre.docx')
for el in list(d.element.body):
    if el.tag!=qn('w:sectPr'): d.element.body.remove(el)
s=d.sections[0]
s.page_width=Cm(21);s.page_height=Cm(29.7)
s.top_margin=s.bottom_margin=Cm(1.5)
s.left_margin=s.right_margin=Cm(1.7)
st=d.styles['Normal'];st.font.name='Arial';st.font.size=Pt(11)
st.paragraph_format.space_after=Pt(1)
st.paragraph_format.line_spacing=1
def p(t,b=False):
    x=d.add_paragraph();r=x.add_run(t);r.bold=b
    return x
def h(t):
    x=p(t,True);x.paragraph_format.space_before=Pt(5);x.paragraph_format.keep_with_next=True
def lines(t):
    for a in t.split('\n'):p(a)
def start(n):
    if n>1:d.add_page_break()
    h(f'PROBLEMA {n}')
def shot(n,width=4.2):
    h('CORRIDA LINDO')
    d.add_paragraph().add_run().add_picture(str(w/f'P{n}.png'),width=Inches(width))
def table(headers,rows):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    for c,x in zip(t.rows[0].cells,headers):c.text=x
    for row in rows:
        for c,x in zip(t.add_row().cells,row):c.text=str(x)
    for row in t.rows:
        for cell in row.cells:
            for par in cell.paragraphs:
                par.paragraph_format.space_after=Pt(1)
                for r in par.runs:r.font.size=Pt(10)
            tcpr=cell._tc.get_or_add_tcPr();bd=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');bd.append(e)
            tcpr.append(bd)
    for c in t.rows[0].cells:
        for r in c.paragraphs[0].runs:r.bold=True
    return t

start(1)
h('I. Identificación de las variables')
lines('Xij = número de estudiantes de minoría del distrito i asignados a la escuela j.\nYij = número de estudiantes de mayoría del distrito i asignados a la escuela j.\ni = 1, 2, 3; j = 1, 2.')
h('III. Identificación de las restricciones')
lines('X11 + X12 = 50                 Y11 + Y12 = 200\nX21 + X22 = 50                 Y21 + Y22 = 250\nX31 + X32 = 100               Y31 + Y32 = 150\n300 ≤ X11+X21+X31+Y11+Y21+Y31 ≤ 500\n300 ≤ X12+X22+X32+Y12+Y22+Y32 ≤ 500\n4X11+4X21+4X31−Y11−Y21−Y31 ≥ 0\n7X11+7X21+7X31−3Y11−3Y21−3Y31 ≤ 0\n4X12+4X22+4X32−Y12−Y22−Y32 ≥ 0\n7X12+7X22+7X32−3Y12−3Y22−3Y32 ≤ 0\nXij, Yij ≥ 0')
h('II. Identificación de la Función Objetivo')
lines('MIN Z = X11+2X12+2X21+X22+X31+X32\n              +Y11+2Y12+2Y21+Y22+Y31+Y32')
h('SOLUCIÓN')
lines('Z ≥ 250(1)+300(1)+250(1) = 800\nX11=50   X12=0   X21=0   X22=50   X31=10   X32=90\nY11=200   Y12=0   Y21=0   Y22=250   Y31=40   Y32=110\nEscuela 1: 300 estudiantes; 60/300 = 20%.\nEscuela 2: 500 estudiantes; 140/500 = 28%.\nZ = 250+300+250 = 800')
h('RESPUESTA')
p('Zmin = 800 estudiante-millas.')
shot(1,4.05)

start(2)
h('I. Identificación de las variables')
lines('Pij = litros del licor i producidos en el mes j.\nVij = litros del licor i vendidos en el mes j.\nIij = litros del licor i en inventario al final del mes j.\ni = 1: Añejo; 2: Frutado; 3: Seco.   j = 1, 2, 3, 4.')
h('III. Identificación de las restricciones')
lines('0.30P1j+0.10P2j+0.20P3j ≤ 6000       j = 1, 2, 3, 4\n0.20P1j+0.60P2j+0.40P3j ≤ 7500       j = 1, 2, 3, 4\n0.50P1j+0.30P2j+0.40P3j ≤ 5600       j = 1, 2, 3, 4\nIi,j−1+Pij−Vij = Iij       i = 1, 2, 3; j = 1, 2, 3, 4\nIi0 = 0       i = 1, 2, 3\nV11≤10000   V12≤12000   V13≤14000   V14≤8000\nV21≤12000   V22≤9000     V23≤12000   V24≤14000\nV31≤11000   V32≤13000   V33≤10000   V34≤8000\nPij, Vij, Iij ≥ 0')
h('II. Identificación de la Función Objetivo')
lines('c1 = 0.30(35)+0.20(25)+0.50(20) = 25.50\nc2 = 0.10(35)+0.60(25)+0.30(20) = 24.50\nc3 = 0.20(35)+0.40(25)+0.40(20) = 25.00\nMAX Z = Σⱼ₌₁⁴(50V1j+55V2j+45V3j−25.5P1j−24.5P2j−25P3j)\n              −0.5Σᵢ₌₁³Σⱼ₌₁⁴ Iij')
h('SOLUCIÓN')
lines('P3j = 0\n0.20P1j+0.60P2j = 7500\n0.50P1j+0.30P2j = 5600\nP1j = 4625;   P2j = 32875/3 = 10958.333…\nI22 = 32875/3−9000 = 5875/3\nI23 = 5875/3+32875/3−12000 = 2750/3\nV24 = 2750/3+32875/3 = 11875\nZ = 24.5(18500)+30.5(131500/3)−0.5(8625/3)\nZ = 1788729.166…')
h('RESPUESTA')
table(['Variable','Mes 1','Mes 2','Mes 3','Mes 4'],[
['P1j = V1j','4625','4625','4625','4625'],
['P2j','32875/3','32875/3','32875/3','32875/3'],
['V2j','32875/3','9000','12000','11875'],
['I2j','0','5875/3','2750/3','0'],
['P3j = V3j','0','0','0','0'],
['I1j = I3j','0','0','0','0']])
p('Zmax = $ 1,788,729.17.')
d.add_page_break();shot(2,4.25)

start(3)
h('I. Identificación de las variables')
lines('Qj, Pj, Rj = toneladas extraídas de Q, P y R, respectivamente,\ny enviadas a procesar en la planta j; j = 1, 2.')
h('III. Identificación de las restricciones')
lines('Q1+Q2 ≤ 1000\nP1+P2 ≤ 2000\nR1+R2 ≤ 3000\nQ1+P1+R1 ≤ 2500\nQ2+P2+R2 ≤ 3500\n0.9Q1+0.9Q2+0.9P1+0.9P2+R1+R2 = 4000\n800 ≤ 0.20Q1+0.20Q2+0.10P1+0.10P2+0.35R1+0.35R2 ≤ 3200\n800 ≤ 0.10Q1+0.10Q2+0.20P1+0.20P2+0.30R1+0.30R2 ≤ 3200\n400 ≤ 0.30Q1+0.30Q2+0.30P1+0.30P2+0.20R1+0.20R2 ≤ 3600\n200 ≤ 0.30Q1+0.30Q2+0.30P1+0.30P2+0.15R1+0.15R2 ≤ 3800\nQj, Pj, Rj ≥ 0')
h('II. Identificación de la Función Objetivo')
lines('Q1: 30+20+15=65       Q2: 30+25+12=67\nP1: 40+22+15=77       P2: 40+18+12=70\nR1: 50+15+15=80       R2: 50+21+12=83\nMIN Z = 65Q1+67Q2+77P1+70P2+80R1+83R2')
h('SOLUCIÓN')
lines('Z ≥ 80(4000)−7(Q1+Q2)−2(P1+P2)\nZ ≥ 320000−7(1000)−2(2000) = 309000\nQ1=1000;   P2=2000;   Q2=P1=R2=0\n0.9(1000)+0.9(2000)+R1 = 4000   ⇒   R1=1300\nA=855;   B=890;   C=1160;   D=1095;   A+B+C+D=4000\nPlanta 1: 2300 ≤ 2500;   Planta 2: 2000 ≤ 3500\nZ = 65(1000)+70(2000)+80(1300) = 309000')
h('RESPUESTA')
p('Q1=1000; P2=2000; R1=1300; Q2=P1=R2=0.   Zmin = $ 309,000.')
shot(3,4.05)

start(4)
h('I. Identificación de las variables')
lines('Xi = número de autobuses que inician en el turno i; i = 1, 2, 3, 4, 5, 6.\nInicio de turno: 1=00:00; 2=04:00; 3=08:00; 4=12:00; 5=16:00; 6=20:00.')
h('III. Identificación de las restricciones')
lines('X6+X1 ≥ 20\nX1+X2 ≥ 40\nX2+X3 ≥ 50\nX3+X4 ≥ 35\nX4+X5 ≥ 60\nX5+X6 ≥ 20\nXi ≥ 0, ENTERAS')
h('II. Identificación de la Función Objetivo')
p('MIN Z = X1+X2+X3+X4+X5+X6')
h('SOLUCIÓN')
lines('(X6+X1)+(X2+X3)+(X4+X5) ≥ 20+50+60\nZ ≥ 130\nX1=0;   X2=40;   X3=10;   X4=60;   X5=0;   X6=20\nZ = 0+40+10+60+0+20 = 130')
h('RESPUESTA')
p('Zmin = 130 autobuses.')
shot(4,4.05)

out=Path(r'C:\Users\ACER\Downloads\EXAMEN_PARCIAL_2026_2_RESUELTO_FINAL.docx')
d.save(out)
print(out)
